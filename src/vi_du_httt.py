"""Ví dụ 2: áp dụng công cụ quản lý rủi ro cho một dự án hệ thống thông tin mẫu.

Dự án mẫu: xây dựng cổng đăng ký dịch vụ trực tuyến. Rủi ro được rút ra từ bài học của dự án
Healthcare.gov (GAO-14-694) và xếp theo 13 nhóm rủi ro của Kassem (2022). Thời lượng ba điểm và
xác suất, tác động của rủi ro là THÔNG SỐ GIẢ ĐỊNH để minh hoạ cách dùng công cụ.

Sinh ra:
  - quan-ly-du-an/vi-du-httt.xlsx : Lịch trình (PERT), Sổ rủi ro (PIM), Kết quả Monte Carlo
  - quan-ly-du-an/vi-du-httt.xml  : lịch trình cho MS Project
  - results/figures/mc_*.png, results/tables/mc_*.csv, results/vi-du-httt.md

Chạy:  python src/vi_du_httt.py
"""
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "quan-ly-du-an"
FIG = ROOT / "results" / "figures"
TAB = ROOT / "results" / "tables"
SO_LAN = 10_000
SEED = 2026


def hien_truc(bd):
    """openpyxl 3.1 mặc định ẩn trục khi mở bằng Excel mới: bật lại trục, chú thích đặt dưới."""
    bd.x_axis.delete = False
    bd.y_axis.delete = False
    if bd.title is not None:
        bd.title.overlay = False
    if bd.legend is not None:
        bd.legend.position = "b"
        bd.legend.overlay = False
    bd.x_axis.title = None                               # tên trục dễ đè lên số; tiêu đề biểu đồ đã đủ rõ
    bd.y_axis.title = None
    return bd


HAN_CHOT = 95          # số ngày làm việc cam kết trong hợp đồng (giả định)

# (mã, công việc, tiền nhiệm, lạc quan, khả dĩ nhất, bi quan) – đơn vị: ngày làm việc
CONG_VIEC = [
    ("T1", "Khảo sát, thu thập yêu cầu", [], 8, 10, 15),
    ("T2", "Phân tích, đặc tả yêu cầu", ["T1"], 8, 10, 16),
    ("T3", "Thiết kế kiến trúc và CSDL", ["T2"], 8, 10, 14),
    ("T4", "Thiết kế giao diện (UI/UX)", ["T2"], 6, 8, 12),
    ("T5", "Phát triển back-end", ["T3"], 20, 25, 35),
    ("T6", "Phát triển front-end", ["T4"], 18, 22, 32),
    ("T7", "Tích hợp với hệ thống hiện có", ["T5"], 8, 10, 18),
    ("T8", "Kiểm thử hệ thống, kiểm thử tải", ["T6", "T7"], 10, 12, 18),
    ("T9", "Kiểm thử chấp nhận (UAT)", ["T8"], 5, 6, 10),
    ("T10", "Đào tạo người dùng", ["T9"], 3, 4, 6),
    ("T11", "Chuyển đổi dữ liệu", ["T7"], 5, 6, 10),
    ("T12", "Triển khai chính thức", ["T10", "T11"], 2, 3, 5),
]

# (mã, nhóm theo Kassem, rủi ro, bài học Healthcare.gov, xác suất, trễ tối thiểu, trễ tối đa, công việc bị ảnh hưởng)
RUI_RO = [
    ("R1", "CL", "Chủ đầu tư thay đổi yêu cầu sau khi chốt đặc tả", "Yêu cầu thay đổi liên tục gây chậm và lãng phí công sức nhà thầu", 0.40, 5, 15, "T5"),
    ("R2", "CL", "Chủ đầu tư chậm phê duyệt kết quả UAT", "Đánh giá sẵn sàng bị dời sát ngày vận hành, thiếu phê duyệt", 0.35, 3, 8, "T9"),
    ("R3", "CO", "Đội phát triển thiếu kinh nghiệm công nghệ mới", "Cách phát triển tăng dần lần đầu được áp dụng", 0.30, 5, 12, "T5"),
    ("R4", "FD", "Ước tính khối lượng front-end quá thấp", "Giao việc khi chưa biết quy mô người dùng", 0.35, 4, 10, "T6"),
    ("R5", "TC", "Hợp đồng không rõ phạm vi, phải đàm phán lại", "Hợp đồng chi phí thực thanh toán, ký trước khi rõ yêu cầu", 0.20, 3, 10, "T2"),
    ("R6", "RM", "Thiếu nhân lực do nhân sự chủ chốt nghỉ việc", "—", 0.25, 5, 10, "T6"),
    ("R7", "MR", "Giám sát nhà thầu lỏng lẻo, lỗi phát hiện muộn", "Không dùng kế hoạch giám sát chất lượng sẵn có", 0.30, 3, 8, "T8"),
    ("R8", "PO", "Thay đổi quy định bảo vệ dữ liệu cá nhân", "—", 0.15, 5, 10, "T3"),
    ("R9", "LP", "Người dùng nội bộ chống đối quy trình mới", "—", 0.30, 2, 5, "T10"),
    ("R10", "SE", "Phát hiện lỗ hổng bảo mật khi kiểm thử", "Vận hành khi chưa kiểm chứng đáp ứng yêu cầu", 0.25, 4, 10, "T8"),
    ("R11", "MR", "Tích hợp hệ thống cũ khó hơn dự kiến", "Cổng trao đổi dữ liệu tăng chi phí gần gấp ba", 0.40, 4, 12, "T7"),
    ("R12", "FM", "Sự cố hạ tầng, thiên tai lúc triển khai", "—", 0.05, 5, 15, "T12"),
]


def muc_xac_suat(p):
    return 1 if p < 0.10 else 2 if p < 0.20 else 3 if p < 0.30 else 4 if p < 0.40 else 5


def muc_tac_dong(tre_tb):
    return 1 if tre_tb < 3 else 2 if tre_tb < 5 else 3 if tre_tb < 7 else 4 if tre_tb < 10 else 5


# ---------------------------------------------------------------- tính lịch
def duong_gang(thoi_gian: dict):
    """Lan truyền thuận: thời điểm kết thúc sớm nhất của từng việc và của cả dự án."""
    xong = {}
    for ma, _, truoc, *_ in CONG_VIEC:
        bat_dau = max((xong[t] for t in truoc), default=0)
        xong[ma] = bat_dau + thoi_gian[ma]
    return xong


def duong_gang_vector(D: pd.DataFrame):
    xong = {}
    for ma, _, truoc, *_ in CONG_VIEC:
        bat_dau = np.max(np.column_stack([xong[t] for t in truoc]), axis=1) if truoc else 0
        xong[ma] = bat_dau + D[ma].values
    return pd.DataFrame(xong)


def duong_tới_han(thoi_gian):
    xong = duong_gang(thoi_gian)
    cuoi = max(xong, key=xong.get)
    duong, hien = [cuoi], cuoi
    truoc_cua = {ma: truoc for ma, _, truoc, *_ in CONG_VIEC}
    while truoc_cua[hien]:
        hien = max(truoc_cua[hien], key=lambda t: xong[t])
        duong.append(hien)
    return list(reversed(duong)), xong[cuoi]


# ---------------------------------------------------------------- Monte Carlo
def monte_carlo():
    rng = np.random.default_rng(SEED)
    D = pd.DataFrame({ma: rng.triangular(o, m, p, SO_LAN) for ma, _, _, o, m, p in CONG_VIEC})
    xay_ra = pd.DataFrame({r[0]: rng.random(SO_LAN) < r[4] for r in RUI_RO})
    D_rr = D.copy()
    for ma, _, _, _, p, tmin, tmax, viec in RUI_RO:
        tre = rng.uniform(tmin, tmax, SO_LAN)
        D_rr[viec] += np.where(xay_ra[ma], tre, 0)
    tong_khong_rr = duong_gang_vector(D).max(axis=1)
    xong = duong_gang_vector(D_rr)
    tong = xong.max(axis=1)

    # chỉ số tới hạn: tỉ lệ số lần mỗi việc nằm trên đường găng
    toi_han = {ma: 0 for ma, *_ in CONG_VIEC}
    for k in range(SO_LAN):
        tg = {ma: D_rr.at[k, ma] for ma, *_ in CONG_VIEC}
        duong, _ = duong_tới_han(tg)
        for ma in duong:
            toi_han[ma] += 1
    toi_han = pd.Series(toi_han) / SO_LAN

    # độ nhạy: trung bình thời gian dự án khi rủi ro xảy ra trừ khi không xảy ra
    nhay = pd.Series({r[0]: tong[xay_ra[r[0]]].mean() - tong[~xay_ra[r[0]]].mean() for r in RUI_RO})
    return tong_khong_rr, tong, toi_han, nhay


# ---------------------------------------------------------------- Excel
FONT = "Arial"
VIEN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def tieu_de(ws, hang, cot, rong):
    for i, (t, w) in enumerate(zip(cot, rong), 1):
        c = ws.cell(hang, i, t)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = VIEN
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hang].height = 30


def o(ws, r, c, v, nhap=False, fmt=None, dam=False, can="center"):
    x = ws.cell(r, c, v)
    x.font = Font(name=FONT, size=10, bold=dam, color="0000FF" if nhap else "000000")
    if nhap:
        x.fill = PatternFill("solid", fgColor="FFF2CC")
    if fmt:
        x.number_format = fmt
    x.alignment = Alignment(horizontal=can, vertical="center", wrap_text=True)
    x.border = VIEN


def ghi_excel(ket_qua):
    tong_khong_rr, tong, toi_han, nhay = ket_qua
    wb = Workbook()

    ws = wb.active
    ws.title = "Lịch trình"
    ws["A1"] = "LỊCH TRÌNH DỰ ÁN CỔNG ĐĂNG KÝ DỊCH VỤ TRỰC TUYẾN (ước lượng ba điểm, PERT)"
    ws["A1"].font = Font(name=FONT, size=12, bold=True, color="1F4E78")
    ws["A2"] = "Đơn vị: ngày làm việc. Ô vàng: thông số giả định để minh hoạ. Kỳ vọng PERT = (O + 4M + P)/6; độ lệch chuẩn = (P − O)/6."
    ws["A2"].font = Font(name=FONT, size=9, italic=True)
    tieu_de(ws, 4, ["Mã", "Công việc", "Tiền nhiệm", "Lạc quan (O)", "Khả dĩ (M)", "Bi quan (P)", "Kỳ vọng PERT",
                    "Độ lệch chuẩn", "Chỉ số tới hạn (MC)"], [6, 34, 11, 10, 10, 10, 11, 11, 12])
    for k, (ma, ten, truoc, o_, m_, p_) in enumerate(CONG_VIEC):
        r = 5 + k
        o(ws, r, 1, ma)
        o(ws, r, 2, ten, can="left")
        o(ws, r, 3, ", ".join(truoc) or "—")
        o(ws, r, 4, o_, nhap=True)
        o(ws, r, 5, m_, nhap=True)
        o(ws, r, 6, p_, nhap=True)
        o(ws, r, 7, f"=(D{r}+4*E{r}+F{r})/6", fmt="0.0")
        o(ws, r, 8, f"=(F{r}-D{r})/6", fmt="0.00")
        o(ws, r, 9, round(float(toi_han[ma]), 3), fmt="0%")
    cuoi = 4 + len(CONG_VIEC)
    duong, dai = duong_tới_han({ma: m for ma, _, _, _, m, _ in CONG_VIEC})
    r = cuoi + 2
    ws.cell(r, 2, "Đường găng theo thời lượng khả dĩ:").font = Font(name=FONT, bold=True)
    ws.cell(r, 3, " → ".join(duong))
    ws.cell(r + 1, 2, "Độ dài đường găng (ngày):").font = Font(name=FONT, bold=True)
    hang_duong = [5 + [c[0] for c in CONG_VIEC].index(ma) for ma in duong]
    ws.cell(r + 1, 3, "=" + "+".join(f"E{h}" for h in hang_duong))
    ws.cell(r + 2, 2, "Kỳ vọng PERT đường găng (ngày):").font = Font(name=FONT, bold=True)
    ws.cell(r + 2, 3, "=" + "+".join(f"G{h}" for h in hang_duong)).number_format = "0.0"
    ws.cell(r + 3, 2, "Độ lệch chuẩn đường găng (ngày):").font = Font(name=FONT, bold=True)
    ws.cell(r + 3, 3, "=SQRT(" + "+".join(f"H{h}^2" for h in hang_duong) + ")").number_format = "0.0"
    ws.cell(r + 4, 2, "Cột I: kết quả Monte Carlo do src/vi_du_httt.py tính.").font = Font(name=FONT, size=9, italic=True)

    ws = wb.create_sheet("Sổ rủi ro")
    ws["A1"] = "SỔ ĐĂNG KÝ RỦI RO – XẾP THEO 13 NHÓM CỦA KASSEM (2022), BÀI HỌC TỪ HEALTHCARE.GOV (GAO-14-694)"
    ws["A1"].font = Font(name=FONT, size=12, bold=True, color="1F4E78")
    ws["A2"] = ("Xác suất và số ngày trễ là giả định. Mức P: <10% →1, <20% →2, <30% →3, <40% →4, còn lại 5. "
                "Mức I theo số ngày trễ trung bình: <3 →1, <5 →2, <7 →3, <10 →4, còn lại 5. Điểm = P × I; Đỏ ≥ 15, Vàng 8–14.")
    ws["A2"].font = Font(name=FONT, size=9, italic=True)
    tieu_de(ws, 4, ["Mã", "Nhóm", "Rủi ro", "Bài học Healthcare.gov", "Xác suất", "Trễ min", "Trễ max", "Trễ TB",
                    "Trễ kỳ vọng", "Mức P", "Mức I", "Điểm", "Vùng", "Việc bị ảnh hưởng", "Độ nhạy MC (ngày)"],
            [6, 7, 34, 34, 9, 8, 8, 8, 9, 7, 7, 7, 8, 10, 11])
    for k, (ma, nhom, rr, bh, p, tmin, tmax, viec) in enumerate(RUI_RO):
        r = 5 + k
        o(ws, r, 1, ma)
        o(ws, r, 2, nhom)
        o(ws, r, 3, rr, can="left")
        o(ws, r, 4, bh, can="left")
        o(ws, r, 5, p, nhap=True, fmt="0%")
        o(ws, r, 6, tmin, nhap=True)
        o(ws, r, 7, tmax, nhap=True)
        o(ws, r, 8, f"=(F{r}+G{r})/2", fmt="0.0")
        o(ws, r, 9, f"=E{r}*H{r}", fmt="0.0")
        o(ws, r, 10, f"=IF(E{r}<0.1,1,IF(E{r}<0.2,2,IF(E{r}<0.3,3,IF(E{r}<0.4,4,5))))")
        o(ws, r, 11, f"=IF(H{r}<3,1,IF(H{r}<5,2,IF(H{r}<7,3,IF(H{r}<10,4,5))))")
        o(ws, r, 12, f"=J{r}*K{r}", dam=True)
        o(ws, r, 13, f'=IF(L{r}>=15,"Đỏ",IF(L{r}>=8,"Vàng","Xanh"))')
        o(ws, r, 14, viec)
        o(ws, r, 15, round(float(nhay[ma]), 2), fmt="0.0")
    cuoi = 4 + len(RUI_RO)
    for mau, ten in [("FFC7CE", "Đỏ"), ("FFEB9C", "Vàng"), ("C6EFCE", "Xanh")]:
        ws.conditional_formatting.add(f"M5:M{cuoi}", CellIsRule(operator="equal", formula=[f'"{ten}"'],
                                      fill=PatternFill("solid", fgColor=mau)))
    o(ws, cuoi + 2, 3, "Tổng trễ kỳ vọng nếu cộng dồn (ngày)", dam=True, can="left")
    o(ws, cuoi + 2, 9, f"=SUM(I5:I{cuoi})", fmt="0.0", dam=True)

    ws = wb.create_sheet("Monte Carlo")
    ws["A1"] = f"KẾT QUẢ MÔ PHỎNG MONTE CARLO ({SO_LAN:,} LẦN)".replace(",", ".")
    ws["A1"].font = Font(name=FONT, size=12, bold=True, color="1F4E78")
    ws["A2"] = "Giá trị do src/vi_du_httt.py tính; thời lượng theo phân phối tam giác (O, M, P), rủi ro xảy ra theo Bernoulli."
    ws["A2"].font = Font(name=FONT, size=9, italic=True)
    tieu_de(ws, 4, ["Chỉ số", "Không tính rủi ro", "Có rủi ro"], [34, 16, 16])
    dong = [("Trung bình (ngày)", tong_khong_rr.mean(), tong.mean()),
            ("Độ lệch chuẩn (ngày)", tong_khong_rr.std(), tong.std()),
            ("P50 (ngày)", tong_khong_rr.quantile(0.5), tong.quantile(0.5)),
            ("P80 (ngày)", tong_khong_rr.quantile(0.8), tong.quantile(0.8)),
            ("P90 (ngày)", tong_khong_rr.quantile(0.9), tong.quantile(0.9)),
            (f"Xác suất xong trong {HAN_CHOT} ngày", (tong_khong_rr <= HAN_CHOT).mean(), (tong <= HAN_CHOT).mean())]
    for k, (ten, a, b) in enumerate(dong):
        r = 5 + k
        fmt = "0%" if "Xác suất" in ten else "0.0"
        o(ws, r, 1, ten, can="left")
        o(ws, r, 2, round(float(a), 3), fmt=fmt)
        o(ws, r, 3, round(float(b), 3), fmt=fmt)
    o(ws, 12, 1, "Dự phòng thời gian đề xuất (P80 − kế hoạch)", can="left", dam=True)
    o(ws, 12, 3, f"=C8-'Lịch trình'!C{len(CONG_VIEC) + 7}", fmt="0.0", dam=True)

    # bảng tần suất (từ Python) + xác suất tích lũy (công thức) để vẽ biểu đồ Excel
    bins = np.arange(int(tong_khong_rr.min()) // 5 * 5, int(tong.max()) // 5 * 5 + 10, 5)
    dem0 = np.histogram(tong_khong_rr, bins)[0]
    dem1 = np.histogram(tong, bins)[0]
    tieu_de_cot = ["Đến ngày", "Số lần (không rủi ro)", "Số lần (có rủi ro)", "Tích lũy (có rủi ro)"]
    for i, (t, w) in enumerate(zip(tieu_de_cot, [10, 14, 14, 14])):
        c = ws.cell(4, 5 + i, t)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(5 + i)].width = w
    for k, (b, a0, a1) in enumerate(zip(bins[1:], dem0, dem1)):
        r = 5 + k
        o(ws, r, 5, int(b))
        o(ws, r, 6, int(a0))
        o(ws, r, 7, int(a1))
        o(ws, r, 8, f"=SUM($G$5:G{r})/SUM($G$5:$G${4 + len(dem1)})", fmt="0%")
    cuoi = 4 + len(dem1)
    cot_ch = BarChart()
    cot_ch.title = "Phân phối thời gian hoàn thành (10.000 lần mô phỏng)"
    cot_ch.y_axis.title = "Số lần"
    cot_ch.x_axis.title = "Thời gian hoàn thành (ngày)"
    cot_ch.add_data(Reference(ws, min_col=6, max_col=7, min_row=4, max_row=cuoi), titles_from_data=True)
    cot_ch.set_categories(Reference(ws, min_col=5, min_row=5, max_row=cuoi))
    duong = LineChart()
    duong.add_data(Reference(ws, min_col=8, min_row=4, max_row=cuoi), titles_from_data=True)
    duong.y_axis.axId = 200
    duong.y_axis.crosses = "max"
    duong.y_axis.delete = False
    duong.y_axis.scaling.min, duong.y_axis.scaling.max = 0, 1
    cot_ch += duong
    cot_ch.height, cot_ch.width = 9, 20
    ws.add_chart(hien_truc(cot_ch), f"A{cuoi + 3}")

    # sổ rủi ro: đóng góp kỳ vọng = xác suất × độ nhạy, và biểu đồ
    rr = wb["Sổ rủi ro"]
    c = rr.cell(4, 16, "Đóng góp kỳ vọng (ngày)")
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    rr.column_dimensions["P"].width = 12
    n_rr = len(RUI_RO)
    for r in range(5, 5 + n_rr):
        o(rr, r, 16, f"=E{r}*O{r}", fmt="0.00", dam=True)
    ngang = BarChart()
    ngang.type = "bar"
    ngang.title = "Đóng góp kỳ vọng của từng rủi ro vào thời gian dự án (ngày)"
    ngang.add_data(Reference(rr, min_col=16, min_row=4, max_row=4 + n_rr), titles_from_data=True)
    ngang.set_categories(Reference(rr, min_col=1, min_row=5, max_row=4 + n_rr))
    ngang.legend = None
    ngang.height, ngang.width = 9, 16
    rr.add_chart(hien_truc(ngang), f"C{n_rr + 9}")
    wb.save(OUT / "vi-du-httt.xlsx")


# ---------------------------------------------------------------- MS Project
def ghi_mspdi():
    uid = {ma: k + 1 for k, (ma, *_) in enumerate(CONG_VIEC)}
    tasks = ["<Task><UID>0</UID><ID>0</ID><Name>Cổng đăng ký dịch vụ trực tuyến</Name>"
             "<OutlineLevel>0</OutlineLevel><Summary>1</Summary></Task>"]
    for ma, ten, truoc, o_, m_, p_ in CONG_VIEC:
        lien_ket = "".join(f"<PredecessorLink><PredecessorUID>{uid[t]}</PredecessorUID><Type>1</Type></PredecessorLink>"
                           for t in truoc)
        tasks.append(
            f"<Task><UID>{uid[ma]}</UID><ID>{uid[ma]}</ID><Name>{escape(ten)}</Name><WBS>{uid[ma]}</WBS>"
            f"<OutlineLevel>1</OutlineLevel><Duration>PT{m_ * 8}H0M0S</Duration><DurationFormat>7</DurationFormat>"
            f"<Notes>{escape(f'{ma}: lạc quan {o_}, khả dĩ {m_}, bi quan {p_} ngày')}</Notes>{lien_ket}</Task>")
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Project xmlns="http://schemas.microsoft.com/project">'
           "<Name>Vi du HTTT - Cong dang ky dich vu truc tuyen</Name>"
           "<Title>Cổng đăng ký dịch vụ trực tuyến</Title>"
           "<StartDate>2027-01-04T08:00:00</StartDate><ScheduleFromStart>1</ScheduleFromStart>"
           f"<Tasks>{''.join(tasks)}</Tasks></Project>")
    (OUT / "vi-du-httt.xml").write_text(xml, encoding="utf-8")


# ---------------------------------------------------------------- hình
def ve_hinh(ket_qua):
    tong_khong_rr, tong, toi_han, nhay = ket_qua
    _, ke_hoach = duong_tới_han({ma: m for ma, _, _, _, m, _ in CONG_VIEC})

    fig, ax1 = plt.subplots(figsize=(9, 5))
    bins = np.arange(int(tong_khong_rr.min()) - 1, int(tong.max()) + 2)
    ax1.hist(tong_khong_rr, bins=bins, alpha=0.5, color="#9ecae1", label="Không tính rủi ro")
    ax1.hist(tong, bins=bins, alpha=0.6, color="#3182bd", label="Có rủi ro")
    ax1.set_xlabel("Thời gian hoàn thành dự án (ngày làm việc)")
    ax1.set_ylabel("Số lần mô phỏng")
    ax2 = ax1.twinx()
    x = np.sort(tong.values)
    ax2.plot(x, np.arange(1, len(x) + 1) / len(x), color="#e6550d", lw=2, label="Xác suất tích lũy (có rủi ro)")
    ax2.set_ylim(0, 1.02)
    ax2.set_ylabel("Xác suất hoàn thành")
    for q, nhan in [(0.5, "P50"), (0.8, "P80")]:
        v = tong.quantile(q)
        ax2.axvline(v, color="#e6550d", ls=":", lw=1)
        ax2.text(v, 0.05, f" {nhan} = {v:.0f}", color="#e6550d", fontsize=8)
    ax1.axvline(ke_hoach, color="black", ls="--", lw=1)
    ax1.text(ke_hoach, ax1.get_ylim()[1] * 0.95, f" Kế hoạch {ke_hoach} ngày", fontsize=8)
    ax1.axvline(HAN_CHOT, color="#c00000", ls="-", lw=1)
    ax1.text(HAN_CHOT, ax1.get_ylim()[1] * 0.85, f" Hạn chót {HAN_CHOT}", fontsize=8, color="#c00000")
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc="upper right", fontsize=8)
    ax1.set_title(f"Phân phối thời gian hoàn thành dự án ({SO_LAN:,} lần mô phỏng)".replace(",", "."))
    fig.tight_layout()
    fig.savefig(FIG / "mc_phan_phoi.png", dpi=200)
    plt.close(fig)

    ten = {r[0]: f"{r[0]} ({r[1]}) {r[2]}" for r in RUI_RO}
    xs = pd.Series({r[0]: r[4] for r in RUI_RO})
    ky_vong = (xs * nhay).sort_values()
    fig, ax = plt.subplots(figsize=(9.5, 5))
    ax.barh([ten[k] for k in ky_vong.index], ky_vong.values,
            color=["#c00000" if v >= 2 else "#f4a582" for v in ky_vong.values])
    for y, k in enumerate(ky_vong.index):
        ax.text(ky_vong[k] + 0.05, y, f"p = {xs[k]:.0%}, nếu xảy ra +{nhay[k]:.1f} ngày", va="center", fontsize=7)
    ax.set_xlim(0, ky_vong.max() * 1.6)
    ax.set_xlabel("Đóng góp kỳ vọng vào thời gian dự án = xác suất × số ngày tăng thêm khi xảy ra")
    ax.set_title("Biểu đồ độ nhạy (tornado): rủi ro nào cần ưu tiên ứng phó")
    ax.tick_params(axis="y", labelsize=8)
    fig.tight_layout()
    fig.savefig(FIG / "mc_do_nhay.png", dpi=200)
    plt.close(fig)


def main():
    OUT.mkdir(exist_ok=True)
    kq = monte_carlo()
    tong_khong_rr, tong, toi_han, nhay = kq
    duong, ke_hoach = duong_tới_han({ma: m for ma, _, _, _, m, _ in CONG_VIEC})
    ghi_excel(kq)
    ghi_mspdi()
    ve_hinh(kq)
    pd.DataFrame({"chi_so_toi_han": toi_han}).to_csv(TAB / "mc_chi_so_toi_han.csv", index_label="viec")
    pd.DataFrame({"do_nhay_ngay": nhay}).to_csv(TAB / "mc_do_nhay.csv", index_label="rui_ro")

    dong = [
        "# Ví dụ 2 – Monte Carlo cho dự án HTTT mẫu",
        "",
        "File này do `src/vi_du_httt.py` sinh ra. Thông số dự án là giả định để minh hoạ.",
        "",
        f"- Đường găng (thời lượng khả dĩ): {' → '.join(duong)} = {ke_hoach} ngày",
        f"- Hạn chót hợp đồng (giả định): {HAN_CHOT} ngày",
        f"- Không tính rủi ro: trung bình {tong_khong_rr.mean():.1f}, P50 {tong_khong_rr.quantile(.5):.1f}, "
        f"P80 {tong_khong_rr.quantile(.8):.1f}, P90 {tong_khong_rr.quantile(.9):.1f}; "
        f"xác suất kịp hạn {(tong_khong_rr <= HAN_CHOT).mean():.1%}",
        f"- Có rủi ro: trung bình {tong.mean():.1f}, P50 {tong.quantile(.5):.1f}, P80 {tong.quantile(.8):.1f}, "
        f"P90 {tong.quantile(.9):.1f}; xác suất kịp hạn {(tong <= HAN_CHOT).mean():.1%}",
        f"- Xác suất xong đúng kế hoạch {ke_hoach} ngày khi có rủi ro: {(tong <= ke_hoach).mean():.1%}",
        "",
        "## Chỉ số tới hạn",
        "",
        toi_han.round(3).to_frame("chi_so").to_markdown(),
        "",
        "## Độ nhạy (ngày tăng thêm khi rủi ro xảy ra)",
        "",
        pd.DataFrame({"neu_xay_ra": nhay, "ky_vong": nhay * pd.Series({r[0]: r[4] for r in RUI_RO})}).sort_values("ky_vong", ascending=False).round(2).to_markdown(),
    ]
    (ROOT / "results" / "vi-du-httt.md").write_text("\n".join(dong), encoding="utf-8")
    print("\n".join(dong))


if __name__ == "__main__":
    main()
