"""Sinh hồ sơ quản lý dự án cho chính việc làm báo cáo giữa kỳ.

Từ một bảng công việc duy nhất (CONG_VIEC), script tạo:
  - quan-ly-du-an/quan-ly-du-an.xlsx : WBS, Gantt, sổ đăng ký rủi ro, đối chiếu barem, nhật ký
  - quan-ly-du-an/ke-hoach.xml       : file MSPDI để mở bằng MS Project (File > Open > .xml)
  - results/figures/gantt_bao_cao.png: biểu đồ Gantt kế hoạch và thực tế cho báo cáo

Chạy:  python src/quan_ly_du_an.py
Cập nhật tiến độ: sửa cột thực tế và % hoàn thành trong CONG_VIEC, sửa NGAY_BAO_CAO, chạy lại.
"""
from datetime import date, datetime, timedelta
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "quan-ly-du-an"
NGAY_BAO_CAO = date(2026, 10, 26)
BAT_DAU, KET_THUC = date(2026, 10, 18), date(2026, 11, 4)


def d(ngay, thang=10):
    return date(2026, thang, ngay)


# (WBS, tên, KH bắt đầu, KH kết thúc, TT bắt đầu, TT kết thúc, % xong, giờ công, sản phẩm, mục barem)
# KH = theo PLAN.md ngày 18/10; TT = theo lịch sử commit
CONG_VIEC = [
    ("1", "Khởi động dự án", None, None, None, None, None, None, "", ""),
    ("1.1", "Tạo repo, kế hoạch, cấu trúc thư mục", d(18), d(18), d(19), d(19), 1.0, 3, "PLAN.md, README.md", "GitHub"),
    ("1.2", "Gửi email xin dữ liệu gốc cho tác giả", d(18), d(18), None, None, 0.0, 1, "docs/email-xin-du-lieu.md", "Tái hiện"),
    ("2", "Đọc hiểu bài báo", None, None, None, None, None, None, "", ""),
    ("2.1", "Ghi chú mục 1–3, bảng thuật ngữ", d(19), d(20), d(20), d(20), 1.0, 6, "docs/notes/01, 02, thuat-ngu", "Nội dung"),
    ("2.2", "Ghi chú kết quả, tổng hợp điểm mâu thuẫn", d(21), d(21), d(20), d(20), 1.0, 4, "docs/notes/03, 04", "Mở rộng"),
    ("2.3", "Lập danh sách tài liệu tham khảo", d(21), d(21), d(20), d(20), 1.0, 2, "docs/tai-lieu-tham-khao.md", "Nội dung"),
    ("3", "Tái hiện thực nghiệm", None, None, None, None, None, None, "", ""),
    ("3.1", "Nhập số liệu các bảng của bài báo", d(22), d(23), d(26), d(26), 1.0, 3, "data/published/*.csv", "Tái hiện"),
    ("3.2", "Kiểm tra lại số liệu đã công bố", d(23), d(23), d(26), d(26), 1.0, 5, "src/kiem_tra_so_lieu.py", "Tái hiện"),
    ("3.3", "Sinh dữ liệu mô phỏng 314 phiếu", d(27), d(27), None, None, 0.0, 4, "data/simulated", "Tái hiện"),
    ("3.4", "Cronbach, RII, PIM trên dữ liệu mô phỏng", d(27), d(27), None, None, 0.0, 4, "results", "Tái hiện"),
    ("3.5", "PLS-SEM, bootstrapping, HTMT", d(28), d(28), None, None, 0.0, 8, "src, results", "Tái hiện"),
    ("4", "Viết báo cáo", None, None, None, None, None, None, "", ""),
    ("4.1", "Dựng khung LaTeX theo mẫu KLTN", d(22), d(22), d(26), d(26), 1.0, 2, "report/main.tex", "Hình thức"),
    ("4.2", "Viết Chương 1 – Tổng quan", d(24), d(24), d(26), d(26), 1.0, 5, "report/chuong/chuong1.tex", "Nội dung"),
    ("4.3", "Vẽ sơ đồ tư duy và sơ đồ mô hình", d(25), d(25), d(26), d(26), 1.0, 3, "report/hinh", "Nội dung"),
    ("4.4", "Viết Chương 2 – Phương pháp và mô hình", d(25), d(26), d(26), d(26), 1.0, 8, "report/chuong/chuong2.tex", "Nội dung"),
    ("4.5", "Viết Chương 3 – Tái hiện thực nghiệm", d(29), d(29), None, None, 0.0, 8, "report/chuong/chuong3.tex", "Tái hiện"),
    ("4.6", "Ví dụ thực tế 1: dự án dầu khí Việt Nam", d(30), d(30), None, None, 0.0, 6, "chuong4.tex", "Mở rộng"),
    ("4.7", "Ví dụ thực tế 2: dự án HTTT (MS Project, Excel)", d(31), d(31), None, None, 0.0, 8, "chuong4.tex", "Mở rộng"),
    ("4.8", "Mô phỏng Monte Carlo, phản biện bài báo", d(1, 11), d(1, 11), None, None, 0.0, 6, "chuong4.tex", "Mở rộng"),
    ("4.9", "Mở đầu, kết luận, rà soát, xuất PDF", d(2, 11), d(2, 11), None, None, 0.0, 6, "bao-cao.pdf", "Hình thức"),
    ("5", "Thuyết trình và nộp bài", None, None, None, None, None, None, "", ""),
    ("5.1", "Làm slide (≤ 15 slide nội dung)", d(3, 11), d(3, 11), None, None, 0.0, 4, "slides/", "Thuyết trình"),
    ("5.2", "Tập thuyết trình, soạn câu hỏi", d(3, 11), d(3, 11), None, None, 0.0, 3, "", "Thuyết trình"),
    ("5.3", "Nộp bài", d(4, 11), d(4, 11), None, None, 0.0, 0, "Mốc", "Hình thức"),
    ("6", "Quản lý dự án", None, None, None, None, None, None, "", ""),
    ("6.1", "Lập WBS, Gantt, sổ rủi ro (Excel, MS Project)", d(26), d(26), d(26), d(26), 1.0, 2, "quan-ly-du-an/", "Mở rộng"),
]

# (mã, mô tả, nguyên nhân, xác suất 1-5, tác động 1-5, chiến lược, hành động, người phụ trách, trạng thái)
RUI_RO = [
    ("R1", "Không có dữ liệu gốc của bài báo", "Tác giả chỉ cung cấp khi có yêu cầu", 5, 4, "Giảm thiểu",
     "Kiểm tra lại số liệu đã công bố + dùng dữ liệu mô phỏng; gửi email xin dữ liệu", "Sinh viên", "Đang xử lý"),
    ("R2", "Lịch sử GitHub ngắn hơn 3 tuần", "Bắt đầu muộn (19/10, hạn 04/11)", 5, 3, "Chấp nhận",
     "Commit đều mỗi ngày; chủ động báo giảng viên", "Sinh viên", "Đã xảy ra"),
    ("R3", "Chậm tiến độ so với kế hoạch", "Đã trễ 4 ngày ở giai đoạn 22–25/10", 4, 4, "Giảm thiểu",
     "Làm bù ngày 26/10; ưu tiên mục có trong barem", "Sinh viên", "Đang xử lý"),
    ("R4", "Không biên dịch được LaTeX trên máy", "Máy chưa cài TeX", 3, 3, "Né tránh",
     "Biên dịch trên Overleaf (XeLaTeX)", "Sinh viên", "Đang theo dõi"),
    ("R5", "Báo cáo dưới 20 trang ruột", "Chương 3–4 chưa viết", 2, 4, "Giảm thiểu",
     "Đếm trang sau mỗi chương; bổ sung phần mở rộng nếu thiếu", "Sinh viên", "Đang theo dõi"),
    ("R6", "Không trả lời được câu hỏi khi thuyết trình", "Kiến thức PLS-SEM mới với sinh viên", 3, 5, "Giảm thiểu",
     "Đọc lại ghi chú, soạn trước 10 câu hỏi và câu trả lời", "Sinh viên", "Đang theo dõi"),
    ("R7", "Thiếu nguồn tin cậy cho ví dụ dầu khí Việt Nam", "Nhiều thông tin chỉ có trên báo mạng", 3, 3, "Giảm thiểu",
     "Dùng bài báo khoa học, kết luận thanh tra, báo cáo kiểm toán", "Sinh viên", "Đang theo dõi"),
    ("R8", "Thiếu phụ lục khai báo công cụ hỗ trợ", "Chưa viết phụ lục B", 3, 5, "Giảm thiểu",
     "Hoàn thiện phụ lục B trước khi xuất PDF (02/11)", "Sinh viên", "Mở"),
    ("R9", "Mất dữ liệu trên máy cá nhân", "Hỏng ổ cứng, xoá nhầm", 1, 5, "Chuyển giao",
     "Push lên GitHub sau mỗi buổi làm", "Sinh viên", "Đang theo dõi"),
]

# (yêu cầu theo slide 52, điểm tối đa, mức hoàn thành hiện tại, nơi đáp ứng)
BAREM = [
    ("Nội dung báo cáo: khái niệm, thành phần, mô hình, cách hoạt động, ứng dụng, ưu/nhược, xu hướng", 2.5, 0.9,
     "Chương 1, Chương 2"),
    ("Tái hiện thực nghiệm phân tích dữ liệu của bài báo", 2.5, 0.4, "Chương 3, src/, results/"),
    ("Khả năng mở rộng, liên hệ thực tế (2 ví dụ thực tế)", 2.0, 0.2, "Chương 4, docs/notes/04"),
    ("Thuyết trình và trả lời câu hỏi", 2.0, 0.0, "slides/"),
    ("Phản biện các đề tài khác", 1.0, 0.0, "docs/phan-bien.md"),
]

NHAT_KY = [
    (d(19), 1, "Khởi tạo repo, kế hoạch, cấu trúc thư mục"),
    (d(20), 10, "Ghi chú đọc bài báo, thuật ngữ, điểm mâu thuẫn, tài liệu tham khảo"),
    (d(26), 11, "Nhập số liệu, script kiểm tra số liệu, khung LaTeX, Chương 1–2, sơ đồ"),
]

# ---------------------------------------------------------------- định dạng
FONT = "Arial"
VIEN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
NEN_TIEU_DE = PatternFill("solid", fgColor="1F4E78")
NEN_TONG = PatternFill("solid", fgColor="DDEBF7")
NEN_NHAP = PatternFill("solid", fgColor="FFF2CC")
CHU_NHAP = Font(name=FONT, size=10, color="0000FF")
CHU = Font(name=FONT, size=10)
CHU_DAM = Font(name=FONT, size=10, bold=True)


def tieu_de(ws, hang, cac_cot, rong):
    for i, (ten, w) in enumerate(zip(cac_cot, rong), start=1):
        c = ws.cell(hang, i, ten)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = NEN_TIEU_DE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = VIEN
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hang].height = 32


def o(ws, hang, cot, gia_tri, nhap=False, fmt=None, dam=False, can="left"):
    c = ws.cell(hang, cot, gia_tri)
    c.font = CHU_NHAP if nhap else (CHU_DAM if dam else CHU)
    if nhap:
        c.fill = NEN_NHAP
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=can, vertical="center", wrap_text=True)
    c.border = VIEN
    return c


def ghi_chu_dau_trang(ws, dong, noi_dung):
    c = ws.cell(dong, 1, noi_dung)
    c.font = Font(name=FONT, size=9, italic=True, color="595959")


# ---------------------------------------------------------------- sheet WBS
def sheet_wbs(wb):
    ws = wb.active
    ws.title = "WBS"
    ws["A1"] = "CẤU TRÚC PHÂN RÃ CÔNG VIỆC (WBS) – BÁO CÁO GIỮA KỲ"
    ws["A1"].font = Font(name=FONT, size=13, bold=True, color="1F4E78")
    ws["A2"] = "Ngày báo cáo:"
    ws["A2"].font = CHU_DAM
    o(ws, 2, 2, NGAY_BAO_CAO, nhap=True, fmt="dd/mm/yyyy")
    ws["B2"].comment = Comment("Ngày cập nhật tình trạng. Sửa ô này mỗi lần cập nhật tiến độ.", "Sinh viên")
    ghi_chu_dau_trang(ws, 3, "Ô nền vàng chữ xanh: dữ liệu nhập tay. Các cột còn lại là công thức. "
                             "KH = kế hoạch theo PLAN.md (18/10); TT = thực tế theo lịch sử commit.")

    cot = ["WBS", "Công việc", "KH bắt đầu", "KH kết thúc", "TT bắt đầu", "TT kết thúc", "Thời lượng KH (ngày)",
           "Chênh lệch kết thúc (ngày)", "% hoàn thành", "Giờ công", "Trạng thái", "Sản phẩm", "Mục barem"]
    tieu_de(ws, 5, cot, [6, 44, 11, 11, 11, 11, 11, 12, 11, 9, 13, 30, 12])
    hang = 6
    tong = {}
    for wbs, ten, kb, kk, tb, tk, pt, gio, sp, barem in CONG_VIEC:
        r = hang
        if kb is None:                       # dòng tổng hợp
            tong[wbs] = r
            for c in range(1, 14):
                o(ws, r, c, None, dam=True).fill = NEN_TONG
            ws.cell(r, 1, wbs)
            ws.cell(r, 2, ten)
        else:
            o(ws, r, 1, wbs, can="center")
            o(ws, r, 2, ten)
            o(ws, r, 3, kb, nhap=True, fmt="dd/mm", can="center")
            o(ws, r, 4, kk, nhap=True, fmt="dd/mm", can="center")
            o(ws, r, 5, tb, nhap=True, fmt="dd/mm", can="center")
            o(ws, r, 6, tk, nhap=True, fmt="dd/mm", can="center")
            o(ws, r, 7, f"=D{r}-C{r}+1", fmt="0", can="center")
            o(ws, r, 8, f'=IF(F{r}="","",F{r}-D{r})', fmt='+0;-0;0', can="center")
            o(ws, r, 9, pt, nhap=True, fmt="0%", can="center")
            o(ws, r, 10, gio, nhap=True, fmt="0", can="center")
            o(ws, r, 11, f'=IF(I{r}>=1,"Hoàn thành",IF(I{r}>0,"Đang làm",'
                         f'IF(D{r}<$B$2,"Trễ hạn","Chưa bắt đầu")))', can="center")
            o(ws, r, 12, sp)
            o(ws, r, 13, barem, can="center")
        hang += 1
    cuoi = hang - 1

    # dòng tổng hợp: lấy min/max ngày, % theo giờ công của các việc con
    ma_tong = list(tong)
    for i, wbs in enumerate(ma_tong):
        r = tong[wbs]
        dau, het = r + 1, (tong[ma_tong[i + 1]] - 1 if i + 1 < len(ma_tong) else cuoi)
        ws.cell(r, 3, f"=MIN(C{dau}:C{het})").number_format = "dd/mm"
        ws.cell(r, 4, f"=MAX(D{dau}:D{het})").number_format = "dd/mm"
        ws.cell(r, 7, f"=D{r}-C{r}+1").number_format = "0"
        ws.cell(r, 9, f"=IFERROR(SUMPRODUCT(I{dau}:I{het},J{dau}:J{het})/SUM(J{dau}:J{het}),0)").number_format = "0%"
        ws.cell(r, 10, f"=SUM(J{dau}:J{het})").number_format = "0"
        for c in (3, 4, 7, 9, 10):
            ws.cell(r, c).alignment = Alignment(horizontal="center")

    # tổng toàn dự án
    r = cuoi + 2
    o(ws, r, 2, "TOÀN DỰ ÁN (tính theo giờ công)", dam=True)
    tu = "+".join(f"I{tong[k]}*J{tong[k]}" for k in ma_tong)
    mau = "+".join(f"J{tong[k]}" for k in ma_tong)
    o(ws, r, 9, f"=({tu})/({mau})", fmt="0.0%", dam=True, can="center")
    o(ws, r + 1, 2, "Số việc trễ hạn", dam=True)
    o(ws, r + 1, 9, f'=COUNTIF(K6:K{cuoi},"Trễ hạn")', fmt="0", dam=True, can="center")
    o(ws, r + 2, 2, "Tổng giờ công kế hoạch", dam=True)
    o(ws, r + 2, 9, "=" + "+".join(f"J{tong[k]}" for k in ma_tong), fmt="0", dam=True, can="center")

    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1")
    ws.add_data_validation(dv)
    dv.add(f"I6:I{cuoi}")
    ws.conditional_formatting.add(f"K6:K{cuoi}", CellIsRule(operator="equal", formula=['"Hoàn thành"'],
                                  fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(f"K6:K{cuoi}", CellIsRule(operator="equal", formula=['"Trễ hạn"'],
                                  fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(f"H6:H{cuoi}", CellIsRule(operator="greaterThan", formula=["0"],
                                  font=Font(name=FONT, color="C00000", bold=True)))
    ws.freeze_panes = "C6"
    return ws, cuoi


# ---------------------------------------------------------------- sheet Gantt
def sheet_gantt(wb, cuoi):
    ws = wb.create_sheet("Gantt")
    ws["A1"] = "BIỂU ĐỒ GANTT – KẾ HOẠCH (□) VÀ THỰC TẾ (■)"
    ws["A1"].font = Font(name=FONT, size=13, bold=True, color="1F4E78")
    ghi_chu_dau_trang(ws, 2, "Tự cập nhật theo sheet WBS. Cột tô viền đỏ là ngày báo cáo (WBS!B2).")

    so_ngay = (KET_THUC - BAT_DAU).days + 1
    tieu_de(ws, 4, ["WBS", "Công việc"] + [""] * so_ngay, [6, 40] + [4.2] * so_ngay)
    for k in range(so_ngay):
        c = ws.cell(4, 3 + k, BAT_DAU + timedelta(days=k))
        c.number_format = "dd"
        c.alignment = Alignment(horizontal="center", text_rotation=90)
    ws.cell(3, 3, "Tháng 10/2026").font = CHU_DAM
    ws.cell(3, 3 + (date(2026, 11, 1) - BAT_DAU).days, "Tháng 11").font = CHU_DAM

    for i in range(6, cuoi + 1):
        r = i - 1
        ws.cell(r, 1, f"=WBS!A{i}").font = CHU
        ws.cell(r, 2, f"=WBS!B{i}").font = CHU
        for k in range(so_ngay):
            cl = get_column_letter(3 + k)
            c = ws.cell(r, 3 + k, f'=IF(WBS!$C{i}="","",IF(AND(WBS!$E{i}<>"",{cl}$4>=WBS!$E{i},{cl}$4<=WBS!$F{i}),"■",'
                                  f'IF(AND({cl}$4>=WBS!$C{i},{cl}$4<=WBS!$D{i}),"□","")))')
            c.alignment = Alignment(horizontal="center")
            c.font = Font(name=FONT, size=9)
            c.border = VIEN
    vung = f"C5:{get_column_letter(2 + so_ngay)}{cuoi - 1}"
    ws.conditional_formatting.add(vung, FormulaRule(formula=['C5="■"'], fill=PatternFill("solid", fgColor="2E75B6"),
                                                    font=Font(color="2E75B6")))
    ws.conditional_formatting.add(vung, FormulaRule(formula=['C5="□"'], fill=PatternFill("solid", fgColor="BDD7EE"),
                                                    font=Font(color="BDD7EE")))
    ws.conditional_formatting.add(f"C4:{get_column_letter(2 + so_ngay)}{cuoi - 1}",
                                  FormulaRule(formula=["C$4=WBS!$B$2"],
                                              border=Border(left=Side(style="medium", color="C00000"),
                                                            right=Side(style="medium", color="C00000"))))
    ws.freeze_panes = "C5"


# ---------------------------------------------------------------- sheet rủi ro
def sheet_rui_ro(wb):
    ws = wb.create_sheet("Rủi ro")
    ws["A1"] = "SỔ ĐĂNG KÝ RỦI RO CỦA DỰ ÁN LÀM BÁO CÁO"
    ws["A1"].font = Font(name=FONT, size=13, bold=True, color="1F4E78")
    ghi_chu_dau_trang(ws, 2, "Xác suất (P) và tác động (I) chấm 1–5. Điểm = P × I. Vùng: Đỏ ≥ 15, Vàng 8–14, Xanh < 8 "
                             "(cùng quy ước với ma trận PIM trong báo cáo).")
    cot = ["Mã", "Rủi ro", "Nguyên nhân", "P", "I", "Điểm", "Vùng", "Chiến lược", "Hành động ứng phó",
           "Phụ trách", "Trạng thái"]
    tieu_de(ws, 4, cot, [6, 34, 30, 5, 5, 7, 8, 12, 44, 11, 13])
    dv_pi = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    dv_cl = DataValidation(type="list", formula1='"Né tránh,Chuyển giao,Giảm thiểu,Chấp nhận,Báo cáo lên trên"')
    ws.add_data_validation(dv_pi)
    ws.add_data_validation(dv_cl)
    for k, (ma, mo_ta, ng, p, i, cl, hd, pt, tt) in enumerate(RUI_RO):
        r = 5 + k
        o(ws, r, 1, ma, can="center")
        o(ws, r, 2, mo_ta)
        o(ws, r, 3, ng)
        o(ws, r, 4, p, nhap=True, can="center")
        o(ws, r, 5, i, nhap=True, can="center")
        o(ws, r, 6, f"=D{r}*E{r}", can="center", dam=True)
        o(ws, r, 7, f'=IF(F{r}>=15,"Đỏ",IF(F{r}>=8,"Vàng","Xanh"))', can="center")
        o(ws, r, 8, cl, nhap=True, can="center")
        o(ws, r, 9, hd)
        o(ws, r, 10, pt, can="center")
        o(ws, r, 11, tt, nhap=True, can="center")
    cuoi = 4 + len(RUI_RO)
    dv_pi.add(f"D5:E{cuoi}")
    dv_cl.add(f"H5:H{cuoi}")
    for mau, ten in [("FFC7CE", "Đỏ"), ("FFEB9C", "Vàng"), ("C6EFCE", "Xanh")]:
        ws.conditional_formatting.add(f"G5:G{cuoi}", CellIsRule(operator="equal", formula=[f'"{ten}"'],
                                      fill=PatternFill("solid", fgColor=mau)))

    r = cuoi + 2
    o(ws, r, 2, "Số rủi ro vùng đỏ", dam=True)
    o(ws, r, 6, f'=COUNTIF(G5:G{cuoi},"Đỏ")', dam=True, can="center")
    o(ws, r + 1, 2, "Số rủi ro vùng vàng", dam=True)
    o(ws, r + 1, 6, f'=COUNTIF(G5:G{cuoi},"Vàng")', dam=True, can="center")
    o(ws, r + 2, 2, "Điểm rủi ro trung bình", dam=True)
    o(ws, r + 2, 6, f"=AVERAGE(F5:F{cuoi})", fmt="0.0", dam=True, can="center")
    ws.freeze_panes = "C5"


# ---------------------------------------------------------------- sheet barem
def sheet_barem(wb):
    ws = wb.create_sheet("Barem")
    ws["A1"] = "ĐỐI CHIẾU BAREM CHẤM ĐIỂM (Lesson 1, slide 52)"
    ws["A1"].font = Font(name=FONT, size=13, bold=True, color="1F4E78")
    ghi_chu_dau_trang(ws, 2, "Mức hoàn thành do sinh viên tự đánh giá (ô vàng). Điểm dự kiến = điểm tối đa × mức hoàn thành.")
    tieu_de(ws, 4, ["Yêu cầu", "Điểm tối đa", "Mức hoàn thành", "Điểm dự kiến", "Đáp ứng ở đâu"], [62, 11, 13, 12, 30])
    for k, (yc, diem, ht, noi) in enumerate(BAREM):
        r = 5 + k
        o(ws, r, 1, yc)
        o(ws, r, 2, diem, fmt="0.0", can="center")
        o(ws, r, 3, ht, nhap=True, fmt="0%", can="center")
        o(ws, r, 4, f"=B{r}*C{r}", fmt="0.00", can="center")
        o(ws, r, 5, noi)
    cuoi = 4 + len(BAREM)
    r = cuoi + 1
    o(ws, r, 1, "TỔNG", dam=True)
    o(ws, r, 2, f"=SUM(B5:B{cuoi})", fmt="0.0", dam=True, can="center")
    o(ws, r, 3, f"=D{r}/B{r}", fmt="0%", dam=True, can="center")
    o(ws, r, 4, f"=SUM(D5:D{cuoi})", fmt="0.00", dam=True, can="center")

    ws.cell(r + 2, 1, "Yêu cầu hình thức (slide 54)").font = CHU_DAM
    hinh_thuc = [("Mẫu KLTN, 1 file PDF, khuyến khích LaTeX", "Đạt"),
                 ("Tối thiểu 20 trang ruột", "Chưa kiểm tra"),
                 ("Chữ trong bảng, hình là tiếng Việt", "Đạt"),
                 ("Gõ lại công thức, đánh số ở lề phải", "Đạt"),
                 ("Danh mục tài liệu tham khảo, trích dẫn trong bài", "Đạt"),
                 ("Khai báo công cụ hỗ trợ đã dùng", "Chưa đạt"),
                 ("Đẩy file phân tích lên GitHub, lịch sử ≥ 3 tuần", "Không đạt (16 ngày)"),
                 ("Không dùng nguồn không tin cậy", "Đạt")]
    for k, (yc, tt) in enumerate(hinh_thuc):
        o(ws, r + 3 + k, 1, yc)
        o(ws, r + 3 + k, 3, tt, nhap=True, can="center")
    vung = f"C{r + 3}:C{r + 2 + len(hinh_thuc)}"
    ws.conditional_formatting.add(vung, CellIsRule(operator="equal", formula=['"Đạt"'],
                                  fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(vung, FormulaRule(formula=[f'LEFT(C{r + 3},5)="Chưa "'],
                                  fill=PatternFill("solid", fgColor="FFEB9C")))
    ws.conditional_formatting.add(vung, FormulaRule(formula=[f'LEFT(C{r + 3},5)="Không"'],
                                  fill=PatternFill("solid", fgColor="FFC7CE")))


# ---------------------------------------------------------------- sheet nhật ký
def sheet_nhat_ky(wb):
    ws = wb.create_sheet("Nhật ký")
    ws["A1"] = "NHẬT KÝ LÀM VIỆC (theo lịch sử commit GitHub)"
    ws["A1"].font = Font(name=FONT, size=13, bold=True, color="1F4E78")
    tieu_de(ws, 3, ["Ngày", "Số commit", "Nội dung chính", "Lũy kế commit"], [12, 10, 70, 13])
    for k, (ngay, so, nd) in enumerate(NHAT_KY):
        r = 4 + k
        o(ws, r, 1, ngay, nhap=True, fmt="dd/mm/yyyy", can="center")
        o(ws, r, 2, so, nhap=True, can="center")
        o(ws, r, 3, nd, nhap=True)
        o(ws, r, 4, f"=SUM($B$4:B{r})", can="center")
    r = 4 + len(NHAT_KY)
    o(ws, r, 1, "Tổng", dam=True)
    o(ws, r, 2, f"=SUM(B4:B{r - 1})", dam=True, can="center")


# ---------------------------------------------------------------- MS Project (MSPDI)
def xml_ngay(ngay, gio):
    return f"{ngay.isoformat()}T{gio}"


def ghi_mspdi():
    gio = lambda h: f"PT{int(h)}H0M0S"
    ngay_lam = "".join(
        f"<WeekDay><DayType>{t}</DayType><DayWorking>1</DayWorking><WorkingTimes>"
        "<WorkingTime><FromTime>08:00:00</FromTime><ToTime>12:00:00</ToTime></WorkingTime>"
        "<WorkingTime><FromTime>13:00:00</FromTime><ToTime>17:00:00</ToTime></WorkingTime>"
        "</WorkingTimes></WeekDay>" for t in range(1, 8))
    tasks = [
        "<Task><UID>0</UID><ID>0</ID><Name>Báo cáo giữa kỳ – Quản lý dự án HTTT</Name><OutlineLevel>0</OutlineLevel>"
        "<Summary>1</Summary></Task>"]
    uid = 0
    for wbs, ten, kb, kk, tb, tk, pt, gio_cong, sp, barem in CONG_VIEC:
        uid += 1
        cap = wbs.count(".") + 1
        if kb is None:
            tasks.append(f"<Task><UID>{uid}</UID><ID>{uid}</ID><Name>{escape(ten)}</Name><WBS>{wbs}</WBS>"
                         f"<OutlineNumber>{wbs}</OutlineNumber><OutlineLevel>{cap}</OutlineLevel><Summary>1</Summary></Task>")
            continue
        bd, kt = (tb, tk) if tb else (kb, kk)
        so_gio = max((kt - bd).days + 1, 1) * 8
        moc = gio_cong == 0
        thuc_te = ""
        if tb:
            thuc_te = (f"<ActualStart>{xml_ngay(tb, '08:00:00')}</ActualStart>"
                       f"<ActualFinish>{xml_ngay(tk, '17:00:00')}</ActualFinish>")
        tasks.append(
            f"<Task><UID>{uid}</UID><ID>{uid}</ID><Name>{escape(ten)}</Name><WBS>{wbs}</WBS>"
            f"<OutlineNumber>{wbs}</OutlineNumber><OutlineLevel>{cap}</OutlineLevel><Summary>0</Summary>"
            f"<Milestone>{1 if moc else 0}</Milestone>"
            f"<Start>{xml_ngay(bd, '08:00:00')}</Start><Finish>{xml_ngay(kt, '08:00:00' if moc else '17:00:00')}</Finish>"
            f"<Duration>{gio(0 if moc else so_gio)}</Duration><DurationFormat>7</DurationFormat>"
            f"<ConstraintType>4</ConstraintType><ConstraintDate>{xml_ngay(bd, '08:00:00')}</ConstraintDate>"
            f"<PercentComplete>{int(pt * 100)}</PercentComplete>{thuc_te}"
            f"<Notes>{escape(f'Sản phẩm: {sp}. Mục barem: {barem}. Giờ công ước tính: {gio_cong}')}</Notes>"
            f"<Baseline><Number>0</Number><Start>{xml_ngay(kb, '08:00:00')}</Start>"
            f"<Finish>{xml_ngay(kk, '17:00:00')}</Finish>"
            f"<Duration>{gio(((kk - kb).days + 1) * 8)}</Duration></Baseline></Task>")
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Project xmlns="http://schemas.microsoft.com/project">'
        "<Name>Bao cao giua ky - Quan ly du an HTTT</Name>"
        "<Title>Báo cáo giữa kỳ – Quản lý dự án HTTT</Title>"
        f"<StartDate>{xml_ngay(BAT_DAU, '08:00:00')}</StartDate>"
        f"<StatusDate>{xml_ngay(NGAY_BAO_CAO, '17:00:00')}</StatusDate>"
        "<ScheduleFromStart>1</ScheduleFromStart><CalendarUID>1</CalendarUID>"
        "<MinutesPerDay>480</MinutesPerDay><MinutesPerWeek>3360</MinutesPerWeek><DaysPerMonth>30</DaysPerMonth>"
        "<Calendars><Calendar><UID>1</UID><Name>Lich 7 ngay</Name><IsBaseCalendar>1</IsBaseCalendar>"
        f"<WeekDays>{ngay_lam}</WeekDays></Calendar></Calendars>"
        f"<Tasks>{''.join(tasks)}</Tasks>"
        "<Resources><Resource><UID>1</UID><ID>1</ID><Name>Sinh viên</Name><Type>1</Type></Resource></Resources>"
        "</Project>")
    (OUT / "ke-hoach.xml").write_text(xml, encoding="utf-8")


# ---------------------------------------------------------------- hình Gantt cho báo cáo
def ve_gantt():
    viec = [v for v in CONG_VIEC if v[2] is not None]
    fig, ax = plt.subplots(figsize=(11, 7.5))
    for k, (wbs, ten, kb, kk, tb, tk, pt, *_ ) in enumerate(viec):
        y = len(viec) - k
        ax.barh(y + 0.18, (kk - kb).days + 1, left=mdates.date2num(kb), height=0.32, color="#bdd7ee",
                edgecolor="#5b9bd5", label="Kế hoạch" if k == 0 else None)
        if tb:
            ax.barh(y - 0.18, (tk - tb).days + 1, left=mdates.date2num(tb), height=0.32, color="#2e75b6",
                    label="Thực tế (đã xong)" if wbs == "1.1" else None)
    ax.axvline(mdates.date2num(NGAY_BAO_CAO + timedelta(days=1)), color="#c00000", ls="--", lw=1.2,
               label=f"Ngày báo cáo {NGAY_BAO_CAO:%d/%m}")
    ax.set_yticks([len(viec) - k for k in range(len(viec))], [f"{v[0]}  {v[1]}" for v in viec], fontsize=8)
    ax.xaxis.set_major_locator(mdates.DayLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d/%m"))
    plt.setp(ax.get_xticklabels(), rotation=90, fontsize=8)
    ax.set_xlim(mdates.date2num(BAT_DAU), mdates.date2num(KET_THUC + timedelta(days=1)))
    ax.grid(axis="x", color="#eee")
    ax.set_title("Kế hoạch và thực tế thực hiện báo cáo giữa kỳ")
    ax.legend(loc="lower left", fontsize=8)
    fig.tight_layout()
    fig.savefig(ROOT / "results" / "figures" / "gantt_bao_cao.png", dpi=200)
    plt.close(fig)


def main():
    OUT.mkdir(exist_ok=True)
    wb = Workbook()
    _, cuoi = sheet_wbs(wb)
    sheet_gantt(wb, cuoi)
    sheet_rui_ro(wb)
    sheet_barem(wb)
    sheet_nhat_ky(wb)
    wb.save(OUT / "quan-ly-du-an.xlsx")
    ghi_mspdi()
    ve_gantt()
    print("Đã tạo:", OUT / "quan-ly-du-an.xlsx", OUT / "ke-hoach.xml", "results/figures/gantt_bao_cao.png")


if __name__ == "__main__":
    main()
