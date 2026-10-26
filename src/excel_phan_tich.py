"""Sinh file Excel kiểm tra số liệu bài báo bằng CÔNG THỨC Excel (không gõ cứng kết quả).

Mỗi sheet ứng với một phép kiểm tra ở mục 3.2 của báo cáo:
  RII, AVE-CR, Fornell-Larcker, Cấu trúc (Q², GoF, R²/f²), Mô phỏng (so sánh với kết quả Python)

Chạy:  python src/excel_phan_tich.py   → quan-ly-du-an/phan-tich-bai-bao.xlsx
Sau đó mở bằng Excel (hoặc tính lại tự động) để xem kết quả công thức.
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as cot

ROOT = Path(__file__).resolve().parents[1]
PUB = ROOT / "data" / "published"
TAB = ROOT / "results" / "tables"
OUT = ROOT / "quan-ly-du-an" / "phan-tich-bai-bao.xlsx"

FONT = "Arial"
VIEN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
DO = PatternFill("solid", fgColor="FFC7CE")
XANH = PatternFill("solid", fgColor="C6EFCE")
NOI_BO = ["CL", "CO", "CN", "FD", "TC", "RM", "MR"]


def hien_truc(bd):
    """openpyxl 3.1 mặc định ẩn trục khi mở bằng Excel mới: bật lại trục, chú thích đặt dưới."""
    bd.x_axis.delete = False
    bd.y_axis.delete = False
    if bd.title is not None:
        bd.title.overlay = False
    if bd.legend is not None:
        bd.legend.position = "b"
        bd.legend.overlay = False
    bd.x_axis.title = None                               # tên trục dễ đè lên số; tiêu đề biểu đồ đã đủ rõ
    bd.y_axis.title = None
    return bd


def tieu_de(ws, hang, ten_cot, rong=None, cot_dau=1):
    for i, t in enumerate(ten_cot):
        c = ws.cell(hang, cot_dau + i, t)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = VIEN
        if rong:
            ws.column_dimensions[cot(cot_dau + i)].width = rong[i]
    ws.row_dimensions[hang].height = 30


def o(ws, r, c, v, nhap=False, fmt=None, dam=False, can="center"):
    x = ws.cell(r, c, v)
    x.font = Font(name=FONT, size=10, bold=dam, color="0000FF" if nhap else "000000")
    if nhap:
        x.fill = PatternFill("solid", fgColor="FFF2CC")
    if fmt:
        x.number_format = fmt
    x.alignment = Alignment(horizontal=can, vertical="center", wrap_text=True)
    x.border = VIEN
    return x


def dau_trang(ws, tieu, ghi_chu):
    ws["A1"] = tieu
    ws["A1"].font = Font(name=FONT, size=12, bold=True, color="1F4E78")
    ws["A2"] = ghi_chu
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")


# ---------------------------------------------------------------- sheet hướng dẫn
def sheet_huong_dan(wb):
    ws = wb.active
    ws.title = "Hướng dẫn"
    dong = [
        "KIỂM TRA LẠI SỐ LIỆU BÀI BÁO KASSEM (2022) BẰNG EXCEL",
        "",
        "Bài báo: Risk Management Assessment in Oil and Gas Construction Projects Using PLS-SEM. Gases 2(2), 33–60.",
        "Dữ liệu gốc không được công bố; file này chỉ dùng các con số đã in trong bài (Bảng 6–14).",
        "",
        "Quy ước màu:",
        "  • Ô nền vàng, chữ xanh: số liệu nhập từ bài báo.",
        "  • Ô chữ đen: công thức Excel, tự tính lại khi sửa số liệu.",
        "  • Nền xanh lá: khớp / đạt.   Nền đỏ: lệch / không đạt.",
        "",
        "Các sheet:",
        "  RII              – RII tổng = RII tác động × RII xác suất; quy đổi sang mức của ma trận PIM (Bảng 6, 7)",
        "  AVE-CR           – Tính AVE, CR từ 56 loading (Bảng 8)",
        "  Fornell-Larcker  – Ma trận Bảng 9 dựng lại; so √AVE với tương quan lớn nhất",
        "  Cấu trúc         – Q² = 1 − SSE/SSO; GoF; kiểm tra R², f² từ β và tương quan (Bảng 11–14)",
        "  Mô phỏng         – So sánh với kết quả PLS-SEM trên dữ liệu mô phỏng (tính bằng Python, src/)",
        "",
        "Sai số làm tròn chấp nhận được: ô B20 (mặc định 0,0015).",
    ]
    for k, t in enumerate(dong, 1):
        ws.cell(k, 1, t).font = Font(name=FONT, size=12 if k == 1 else 10, bold=k == 1,
                                     color="1F4E78" if k == 1 else "000000")
    ws["A20"] = "Sai số chấp nhận:"
    ws["A20"].font = Font(name=FONT, bold=True)
    o(ws, 20, 2, 0.0015, nhap=True, fmt="0.0000")
    ws.column_dimensions["A"].width = 20


SAI_SO = "'Hướng dẫn'!$B$20"


# ---------------------------------------------------------------- RII
def sheet_rii(wb):
    ws = wb.create_sheet("RII")
    dau_trang(ws, "RII TỔNG VÀ MỨC MA TRẬN XÁC SUẤT – TÁC ĐỘNG (Bảng 6, 7)",
              "Mức quy đổi theo thang chia đều: <0,2 VL; <0,4 L; <0,6 M; <0,8 H; còn lại VH (giả định, bài báo không nêu).")
    b7 = pd.read_csv(PUB / "bang7_rii_top10.csv")
    b6 = pd.read_csv(PUB / "bang6_pim_top10.csv").set_index("ma")
    tieu_de(ws, 4, ["Mã", "Rủi ro", "RII tác động", "RII xác suất", "RII tổng (bài báo)", "RII tổng tính lại",
                    "Chênh lệch", "Kết quả", "Hạng bài báo", "Hạng tính lại", "Mức tác động", "Mức xác suất",
                    "Bảng 6: tác động", "Bảng 6: xác suất", "PIM khớp RII?"],
            [7, 32, 10, 10, 11, 11, 10, 9, 9, 9, 9, 9, 10, 10, 11])
    muc = lambda c: f'=IF({c}<0.2,"VL",IF({c}<0.4,"L",IF({c}<0.6,"M",IF({c}<0.8,"H","VH"))))'
    n = len(b7)
    for k, r_ in b7.iterrows():
        r = 5 + k
        o(ws, r, 1, r_.ma)
        o(ws, r, 2, r_.rui_ro, can="left")
        o(ws, r, 3, r_.rii_tac_dong, nhap=True, fmt="0.000")
        o(ws, r, 4, r_.rii_xac_suat, nhap=True, fmt="0.000")
        o(ws, r, 5, r_.rii_tong, nhap=True, fmt="0.000")
        o(ws, r, 6, f"=C{r}*D{r}", fmt="0.0000")
        o(ws, r, 7, f"=F{r}-E{r}", fmt="+0.0000;-0.0000;0")
        o(ws, r, 8, f'=IF(ABS(G{r})<={SAI_SO},"Khớp","Lệch")')
        o(ws, r, 9, int(r_.hang_tong), nhap=True)
        o(ws, r, 10, f"=RANK(F{r},$F$5:$F${4 + n})")
        o(ws, r, 11, muc(f"C{r}"))
        o(ws, r, 12, muc(f"D{r}"))
        if r_.ma in b6.index:
            o(ws, r, 13, b6.at[r_.ma, "tac_dong"], nhap=True)
            o(ws, r, 14, b6.at[r_.ma, "xac_suat"], nhap=True)
            o(ws, r, 15, f'=IF(AND(K{r}=M{r},L{r}=N{r}),"Có","Không")')
        else:
            for c in (13, 14):
                o(ws, r, c, "—")
            o(ws, r, 15, "Không có trong Bảng 6")
    cuoi = 4 + n
    ws.conditional_formatting.add(f"H5:H{cuoi}", CellIsRule(operator="equal", formula=['"Khớp"'], fill=XANH))
    ws.conditional_formatting.add(f"H5:H{cuoi}", CellIsRule(operator="equal", formula=['"Lệch"'], fill=DO))
    ws.conditional_formatting.add(f"O5:O{cuoi}", CellIsRule(operator="equal", formula=['"Có"'], fill=XANH))
    ws.conditional_formatting.add(f"O5:O{cuoi}", CellIsRule(operator="equal", formula=['"Không"'], fill=DO))
    r = cuoi + 2
    o(ws, r, 2, "Số dòng RII tổng khớp", dam=True, can="left")
    o(ws, r, 8, f'=COUNTIF(H5:H{cuoi},"Khớp")&"/"&COUNTA(H5:H{cuoi})', dam=True)
    o(ws, r + 1, 2, "Số rủi ro có mức PIM khớp RII", dam=True, can="left")
    o(ws, r + 1, 15, f'=COUNTIF(O5:O{cuoi},"Có")&"/"&(COUNTIF(O5:O{cuoi},"Có")+COUNTIF(O5:O{cuoi},"Không"))', dam=True)

    ch = BarChart()
    ch.type = "bar"
    ch.title = "RII tác động và RII xác suất của top 10 rủi ro"
    ch.y_axis.title = "RII"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 1
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=4, max_row=cuoi), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=cuoi))
    ch.height, ch.width = 9, 16
    ws.add_chart(hien_truc(ch), f"B{cuoi + 5}")
    ws.freeze_panes = "C5"


# ---------------------------------------------------------------- AVE, CR
def sheet_ave_cr(wb):
    ws = wb.create_sheet("AVE-CR")
    dau_trang(ws, "AVE VÀ CR TÍNH LẠI TỪ LOADING (Bảng 8)",
              "AVE = trung bình λ²;  CR = (Σλ)² / ((Σλ)² + Σ(1 − λ²)).  Ngưỡng: AVE ≥ 0,5; CR ≥ 0,7.")
    ld = pd.read_csv(PUB / "bang8_loading.csv")
    cs = pd.read_csv(PUB / "bang8_chi_so_nhom.csv").set_index("nhom")
    tieu_de(ws, 4, ["Nhóm", "Câu", "Loading λ", "λ²"], [7, 7, 10, 9])
    for k, r_ in ld.iterrows():
        r = 5 + k
        o(ws, r, 1, r_.nhom)
        o(ws, r, 2, r_.cau)
        o(ws, r, 3, r_.loading, nhap=True, fmt="0.000")
        o(ws, r, 4, f"=C{r}^2", fmt="0.000")
    cuoi = 4 + len(ld)
    rg = lambda c: f"${c}$5:${c}${cuoi}"
    ws.conditional_formatting.add(f"C5:C{cuoi}", CellIsRule(operator="lessThan", formula=["0.708"], fill=DO))

    tieu_de(ws, 4, ["Nhóm", "Số câu", "AVE tính lại", "AVE bài báo", "Kết quả AVE", "CR tính lại", "CR bài báo",
                    "Kết quả CR", "AVE ≥ 0,5?"], [8, 7, 10, 10, 10, 10, 10, 10, 10], cot_dau=6)
    nhom = list(dict.fromkeys(ld.nhom))
    for k, g in enumerate(nhom):
        r = 5 + k
        o(ws, r, 6, g)
        o(ws, r, 7, f'=COUNTIF({rg("A")},F{r})')
        o(ws, r, 8, f'=SUMIF({rg("A")},F{r},{rg("D")})/G{r}', fmt="0.000")
        o(ws, r, 9, cs.at[g, "ave"], nhap=True, fmt="0.000")
        o(ws, r, 10, f'=IF(ABS(H{r}-I{r})<={SAI_SO},"Khớp","Lệch")')
        s = f'SUMIF({rg("A")},F{r},{rg("C")})'
        o(ws, r, 11, f'={s}^2/({s}^2+G{r}-SUMIF({rg("A")},F{r},{rg("D")}))', fmt="0.000")
        o(ws, r, 12, cs.at[g, "cr"], nhap=True, fmt="0.000")
        o(ws, r, 13, f'=IF(ABS(K{r}-L{r})<={SAI_SO},"Khớp","Lệch")')
        o(ws, r, 14, f'=IF(H{r}>=0.5,"Đạt","Không đạt")')
    hc = 4 + len(nhom)
    for vung in (f"J5:J{hc}", f"M5:M{hc}"):
        ws.conditional_formatting.add(vung, CellIsRule(operator="equal", formula=['"Khớp"'], fill=XANH))
        ws.conditional_formatting.add(vung, CellIsRule(operator="equal", formula=['"Lệch"'], fill=DO))
    ws.conditional_formatting.add(f"N5:N{hc}", CellIsRule(operator="equal", formula=['"Đạt"'], fill=XANH))
    ws.conditional_formatting.add(f"N5:N{hc}", CellIsRule(operator="equal", formula=['"Không đạt"'], fill=DO))
    o(ws, hc + 2, 6, "Ghi chú: nhóm EF tính với cả 5 câu; tác giả nói đã bỏ EF3 nhưng vẫn in đủ 5 loading.",
      can="left").font = Font(name=FONT, size=9, italic=True)

    ch = BarChart()
    ch.title = "AVE tính lại và AVE bài báo công bố"
    ch.y_axis.title = "AVE"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0.4, 0.9
    ch.add_data(Reference(ws, min_col=8, max_col=9, min_row=4, max_row=hc), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=6, min_row=5, max_row=hc))
    ch.height, ch.width = 8, 18
    ws.add_chart(hien_truc(ch), f"F{hc + 4}")


# ---------------------------------------------------------------- Fornell–Larcker
def sheet_fornell_larcker(wb):
    ws = wb.create_sheet("Fornell-Larcker")
    dau_trang(ws, "TIÊU CHÍ FORNELL–LARCKER (Bảng 9, dựng lại từ các đường chéo trong PDF)",
              "Tam giác dưới: số liệu bài báo. Tam giác trên: công thức đối xứng. Ô đỏ: tương quan lớn hơn √AVE của "
              "một trong hai nhóm (vi phạm).")
    fl = pd.read_csv(PUB / "bang9_fornell_larcker.csv", index_col=0)
    ten = list(fl.index)
    n = len(ten)
    r0, c0 = 5, 3                        # ô đầu ma trận C5
    tieu_de(ws, 4, ["Nhóm", "Loại"] + ten + ["√AVE", "Tương quan lớn nhất", "Với nhóm", "Chênh lệch", "Kết quả"],
            [7, 10] + [6.5] * n + [8, 10, 9, 10, 11])
    for i, a in enumerate(ten):
        r = r0 + i
        o(ws, r, 1, a, dam=True)
        o(ws, r, 2, "Nội bộ" if a in NOI_BO else "Bên ngoài")
        for j in range(n):
            c = c0 + j
            if j < i:
                o(ws, r, c, float(fl.iloc[i, j]), nhap=True, fmt="0.000")
            elif j > i:
                o(ws, r, c, f"={cot(c0 + i)}{r0 + j}", fmt="0.000")
            else:
                o(ws, r, c, None).fill = PatternFill("solid", fgColor="D9D9D9")
        hang = f"{cot(c0)}{r}:{cot(c0 + n - 1)}{r}"
        cs = c0 + n
        o(ws, r, cs, float(fl.iloc[i, i]), nhap=True, fmt="0.000")
        o(ws, r, cs + 1, f"=MAX({hang})", fmt="0.000")
        o(ws, r, cs + 2, f"=INDEX(${cot(c0)}$4:${cot(c0 + n - 1)}$4,MATCH({cot(cs + 1)}{r},{hang},0))")
        o(ws, r, cs + 3, f"={cot(cs)}{r}-{cot(cs + 1)}{r}", fmt="+0.000;-0.000")
        o(ws, r, cs + 4, f'=IF({cot(cs + 3)}{r}>0,"Đạt","KHÔNG ĐẠT")', dam=True)
    rc = r0 + n
    o(ws, rc, 1, "√AVE", dam=True)
    o(ws, rc + 1, 1, "Loại", dam=True)
    cs = c0 + n
    for j in range(n):
        o(ws, rc, c0 + j, f"={cot(cs)}{r0 + j}", fmt="0.000")
        o(ws, rc + 1, c0 + j, f"=B{r0 + j}")
    ma_tran = f"{cot(c0)}{r0}:{cot(c0 + n - 1)}{rc - 1}"
    ws.conditional_formatting.add(ma_tran, ColorScaleRule(start_type="num", start_value=0.5, start_color="FFFFFF",
                                                          end_type="num", end_value=0.9, end_color="F4B183"))
    ws.conditional_formatting.add(ma_tran, FormulaRule(
        formula=[f'AND({cot(c0)}{r0}<>"",{cot(c0)}{r0}>MIN(${cot(cs)}{r0},{cot(c0)}${rc}))'],
        fill=PatternFill("solid", fgColor="FF0000"), font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
    kq = f"{cot(cs + 4)}{r0}:{cot(cs + 4)}{rc - 1}"
    ws.conditional_formatting.add(kq, CellIsRule(operator="equal", formula=['"Đạt"'], fill=XANH))
    ws.conditional_formatting.add(kq, CellIsRule(operator="equal", formula=['"KHÔNG ĐẠT"'], fill=DO))

    r = rc + 3
    o(ws, r, 1, "Số nhóm không đạt:", dam=True, can="left")
    o(ws, r, 4, f'=COUNTIF({kq},"KHÔNG ĐẠT")', dam=True)
    o(ws, r + 1, 1, "Tương quan TB:", dam=True, can="left")
    o(ws, r + 1, 4, f"=AVERAGE({ma_tran})", fmt="0.000", dam=True)
    o(ws, r + 2, 1, "Số cặp > 0,80:", dam=True, can="left")
    o(ws, r + 2, 4, f'=COUNTIF({ma_tran},">0.8")/2', dam=True)
    ws.freeze_panes = "C5"
    return {"r0": r0, "c0": c0, "n": n, "rc": rc}


# ---------------------------------------------------------------- Cấu trúc
def sheet_cau_truc(wb, vt):
    ws = wb.create_sheet("Cấu trúc")
    dau_trang(ws, "MÔ HÌNH CẤU TRÚC: Q², GoF, KIỂM TRA R² VÀ f² (Bảng 11–14)",
              "Số phiếu n = 314. SSO = n × số câu của biến, nên SSO/n cho biết dòng nào của Bảng 14 ứng với biến nào.")
    for c, w in zip("ABCDEFG", [34, 11, 11, 11, 11, 13, 22]):
        ws.column_dimensions[c].width = w
    ws["A4"] = "Số phiếu n"
    o(ws, 4, 2, 314, nhap=True)

    tieu_de(ws, 6, ["Bảng 14", "SSO", "SSE", "Q² bài báo", "Q² tính lại", "Số câu = SSO/n", "Suy ra là biến"])
    q = pd.read_csv(PUB / "bang14_q2.csv")
    for k, r_ in q.iterrows():
        r = 7 + k
        o(ws, r, 1, f"Dòng {k + 1}")
        o(ws, r, 2, r_.sso, nhap=True, fmt="0.00")
        o(ws, r, 3, r_.sse, nhap=True, fmt="0.000")
        o(ws, r, 4, r_.q2, nhap=True, fmt="0.000")
        o(ws, r, 5, f"=1-C{r}/B{r}", fmt="0.000")
        o(ws, r, 6, f"=B{r}/$B$4", fmt="0")
        o(ws, r, 7, f'=IF(F{r}=4,"EF (4 câu, đã bỏ EF3)",IF(F{r}=20,"Rủi ro bên ngoài",IF(F{r}=31,"Rủi ro nội bộ","?")))')

    r = 12
    tieu_de(ws, r, ["GoF = √(R² × AVE trung bình)", "R²", "AVE TB", "GoF tính", "GoF bài báo", "", ""])
    o(ws, r + 1, 1, "Theo số bài báo dùng", can="left")
    o(ws, r + 1, 2, 0.720, nhap=True, fmt="0.000")
    o(ws, r + 1, 3, 0.566, nhap=True, fmt="0.000")
    o(ws, r + 1, 4, f"=SQRT(B{r + 1}*C{r + 1})", fmt="0.000")
    o(ws, r + 1, 5, 0.638, nhap=True, fmt="0.000")
    o(ws, r + 2, 1, "AVE TB của 13 nhóm (sheet AVE-CR)", can="left")
    o(ws, r + 2, 2, f"=B{r + 1}", fmt="0.000")
    o(ws, r + 2, 3, "=AVERAGE('AVE-CR'!I5:I17)", fmt="0.000")
    o(ws, r + 2, 4, f"=SQRT(B{r + 2}*C{r + 2})", fmt="0.000")
    o(ws, r + 2, 5, "—")

    r = 17
    tieu_de(ws, r, ["Kiểm tra R², f² từ β và Bảng 9", "Giá trị", "Bài báo", "Chênh lệch", "", "", ""])
    f = "'Fornell-Larcker'!"
    c0, r0, n, rc = vt["c0"], vt["r0"], vt["n"], vt["rc"]
    M = f"{f}${cot(c0)}${r0}:${cot(c0 + n - 1)}${rc - 1}"
    loai_hang = f"{f}$B${r0}:$B${rc - 1}"
    loai_cot = f"{f}${cot(c0)}${rc + 1}:${cot(c0 + n - 1)}${rc + 1}"
    tong = lambda a, b: f'SUMPRODUCT(({loai_hang}="{a}")*({loai_cot}="{b}")*{M})'
    dong = [
        ("β nội bộ → EF", 0.200, None, True),
        ("β bên ngoài → EF", 0.673, None, True),
        ("Tổng tương quan khối nội bộ (cả đường chéo = 1)", f"={tong('Nội bộ', 'Nội bộ')}+7", None, False),
        ("Tổng tương quan khối bên ngoài (cả đường chéo = 1)", f"={tong('Bên ngoài', 'Bên ngoài')}+6", None, False),
        ("Tổng tương quan chéo nội bộ × bên ngoài", f"={tong('Nội bộ', 'Bên ngoài')}", None, False),
        ("r = tương quan giữa TB nhóm nội bộ và TB nhóm bên ngoài", "=B22/SQRT(B20*B21)", None, False),
        ("R² = β₁² + β₂² + 2β₁β₂r", "=B18^2+B19^2+2*B18*B19*B23", 0.720, False),
        ("f² nội bộ = (R² − (β₂ + β₁r)²) / (1 − R²)", "=(B24-(B19+B18*B23)^2)/(1-B24)", 0.042, False),
        ("f² bên ngoài = (R² − (β₁ + β₂r)²) / (1 − R²)", "=(B24-(B18+B19*B23)^2)/(1-B24)", 0.471, False),
    ]
    for k, (ten, gt, bb, nhap) in enumerate(dong):
        rr = r + 1 + k
        o(ws, rr, 1, ten, can="left")
        o(ws, rr, 2, gt, nhap=nhap, fmt="0.000")
        if bb is not None:
            o(ws, rr, 3, bb, nhap=True, fmt="0.000")
            o(ws, rr, 4, f"=B{rr}-C{rr}", fmt="+0.000;-0.000")
    o(ws, 28, 1, "r là xấp xỉ (trung bình đều các nhóm); bài báo dùng trọng số PLS không công bố.", can="left").font = \
        Font(name=FONT, size=9, italic=True)


# ---------------------------------------------------------------- Mô phỏng
def sheet_mo_phong(wb):
    ws = wb.create_sheet("Mô phỏng")
    dau_trang(ws, "SO SÁNH VỚI KẾT QUẢ PLS-SEM TRÊN DỮ LIỆU MÔ PHỎNG",
              "Cột 'Mô phỏng' là kết quả do src/phan_tich_mo_phong.py tính (bootstrapping 5000 lần), "
              "đọc từ results/tables/mp_tong_hieu_ung.csv.")
    th = pd.read_csv(TAB / "mp_tong_hieu_ung.csv", index_col=0)
    tieu_de(ws, 4, ["Nhóm", "β bài báo", "β mô phỏng", "SE bài báo", "SE mô phỏng", "t bài báo", "t mô phỏng",
                    "Tỷ lệ SE mô phỏng / bài báo"], [8, 11, 11, 11, 11, 11, 11, 14])
    for k, (g, r_) in enumerate(th.iterrows()):
        r = 5 + k
        o(ws, r, 1, g)
        o(ws, r, 2, r_.beta_bai_bao, nhap=True, fmt="0.000")
        o(ws, r, 3, round(r_.beta, 4), fmt="0.000")
        o(ws, r, 4, f"=B{r}/F{r}", fmt="0.0000")
        o(ws, r, 5, round(r_.se, 4), fmt="0.0000")
        o(ws, r, 6, r_.t_bai_bao, nhap=True, fmt="0.0")
        o(ws, r, 7, f"=C{r}/E{r}", fmt="0.0")
        o(ws, r, 8, f"=E{r}/D{r}", fmt="0.0")
    cuoi = 4 + len(th)
    o(ws, cuoi + 1, 1, "TB", dam=True)
    for c in "BCDEFGH":
        o(ws, cuoi + 1, " ABCDEFGH".index(c), f"=AVERAGE({c}5:{c}{cuoi})", fmt="0.000", dam=True)
    ch = BarChart()
    ch.title = "Hệ số 13 nhóm rủi ro: bài báo và mô phỏng (tổng hiệu ứng)"
    ch.y_axis.title = "Hệ số"
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=4, max_row=cuoi), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=cuoi))
    ch.height, ch.width = 8, 18
    ws.add_chart(hien_truc(ch), f"A{cuoi + 4}")


def main():
    wb = Workbook()
    sheet_huong_dan(wb)
    sheet_rii(wb)
    sheet_ave_cr(wb)
    vt = sheet_fornell_larcker(wb)
    sheet_cau_truc(wb, vt)
    sheet_mo_phong(wb)
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("Đã tạo", OUT)


if __name__ == "__main__":
    main()
